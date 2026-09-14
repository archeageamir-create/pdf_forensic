from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.models import AnalysisResult, Finding, ResearchProject
from app.reporting.safety import atomic_report_path, validate_report_destination
from app.resources.labels import (
    category_label,
    display_value,
    parameter_description,
    parameter_label,
    source_label,
)


BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
MUTED = RGBColor(0x55, 0x55, 0x55)


class DocxReportExporter:
    """Formal technical report using standard_business_brief + memo_masthead."""

    def __init__(self, template: Path | None = None) -> None:
        self.template = template

    def export(
        self,
        result: AnalysisResult,
        destination: Path,
        project: ResearchProject | None = None,
    ) -> Path:
        validate_report_destination(result, destination)
        destination.parent.mkdir(parents=True, exist_ok=True)
        document = Document(str(self.template)) if self.template and self.template.is_file() else Document()
        self._configure(document)
        self._masthead(document, result, project)
        self._body(document, result, project)
        with atomic_report_path(result, destination) as temporary:
            document.save(temporary)
        return destination

    def _configure(self, document: DocumentType) -> None:
        section = document.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        styles = document.styles
        normal = styles["Normal"]
        _font(normal.font, "Calibri", 11)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.10
        for name, size, color, before, after in (
            ("Heading 1", 16, BLUE, 16, 8),
            ("Heading 2", 13, BLUE, 12, 6),
            ("Heading 3", 12, DARK_BLUE, 8, 4),
        ):
            style = styles[name]
            _font(style.font, "Calibri", size, color=color, bold=True)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True

        header = section.header.paragraphs[0]
        header.text = "Document Forensic Inspector | Технический отчёт"
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _font(header.runs[0].font, "Calibri", 9, color=MUTED)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = footer.add_run("Страница ")
        _font(run.font, "Calibri", 9, color=MUTED)
        _add_page_field(footer)

    def _masthead(
        self, document: DocumentType, result: AnalysisResult, project: ResearchProject | None
    ) -> None:
        title = document.add_paragraph()
        title.paragraph_format.space_after = Pt(4)
        run = title.add_run("ТЕХНИЧЕСКИЙ ОТЧЁТ")
        _font(run.font, "Calibri", 23, bold=True)
        subtitle = document.add_paragraph()
        subtitle.paragraph_format.space_after = Pt(16)
        run = subtitle.add_run("Исследование метаданных и структуры файла")
        _font(run.font, "Calibri", 14, color=MUTED)

        authorship = document.add_paragraph()
        authorship.paragraph_format.space_after = Pt(12)
        run = authorship.add_run(
            "Оболочка, интеграция и сборка: Amirkhan Smagulov. Сторонние "
            "open-source инструменты сохраняют своё авторство и лицензии."
        )
        _font(run.font, "Calibri", 9.5, color=MUTED)

        rows = [
            ("Проект", project.title if project else "Быстрое исследование"),
            ("Номер", project.case_number if project else "Не указан"),
            ("Объект", result.evidence.original_name),
            ("Формат", result.evidence.detected_format),
            ("Дата анализа", result.analyzed_at),
            ("Статус", "Частично" if result.incomplete else "Завершён"),
        ]
        self._key_value_table(document, rows)
        note = document.add_paragraph()
        note.paragraph_format.space_before = Pt(10)
        note.paragraph_format.space_after = Pt(10)
        run = note.add_run(
            "Автоматические наблюдения требуют экспертной оценки. Различия метаданных "
            "сами по себе не являются доказательством подделки."
        )
        _font(run.font, "Calibri", 10, color=RGBColor(0x7A, 0x5A, 0x00), bold=True)

    def _body(
        self, document: DocumentType, result: AnalysisResult, project: ResearchProject | None
    ) -> None:
        document.add_heading("1. Сводка", level=1)
        for item in result.summary:
            document.add_paragraph(_xml_safe(item), style="List Bullet")

        document.add_heading("2. Идентификация объекта", level=1)
        evidence = result.evidence
        self._key_value_table(
            document,
            [
                ("Оригинал", evidence.original_path),
                ("Рабочая копия", evidence.working_copy_path),
                ("MIME", evidence.mime_type),
                ("Размер", f"{evidence.size_bytes:,} байт"),
                ("SHA-256", evidence.sha256),
                ("SHA-1", evidence.sha1),
                ("MD5", evidence.md5),
                ("Режим", display_value(evidence.analysis_mode.value)),
            ],
        )

        document.add_heading("3. Использованные методы", level=1)
        self._tool_table(document, result)

        document.add_heading("4. Область исследования", level=1)
        scope = result.artifacts.get("analysis_scope", {})
        if isinstance(scope, dict):
            document.add_paragraph(_xml_safe(str(scope.get("coverage", "Область проверки не записана в этой версии результата."))))
        document.add_paragraph("Полнота относится к выбранным методам. Неподдерживаемые свойства и отсутствующие инструменты не считаются выполненной проверкой.")
        if result.artifacts.get("pdf_provenance"):
            document.add_heading("Авторство и история правок PDF", level=2)
            self._provenance_table(document, result)

        grouped: dict[str, list[Finding]] = defaultdict(list)
        for finding in result.findings:
            if finding.include_in_report:
                grouped[finding.category].append(finding)
        document.add_heading("5. Результаты исследования", level=1)
        for category, findings in grouped.items():
            document.add_heading(category_label(category), level=2)
            self._finding_table(document, findings)

        document.add_heading("6. Хронология", level=1)
        timeline = result.artifacts.get("timeline", [])
        if isinstance(timeline, list) and timeline:
            table = document.add_table(rows=1, cols=4)
            headers = ("Дата", "Источник", "Событие", "Исходная строка")
            for cell, text in zip(table.rows[0].cells, headers):
                cell.text = text
            for item in timeline:
                if not isinstance(item, dict):
                    continue
                cells = table.add_row().cells
                cells[0].text = _xml_safe(str(item.get("normalized") or item.get("raw", "")) + ("\nЧасовой пояс не указан" if not item.get("timezone_present") else ""))
                cells[1].text = _xml_safe(source_label(str(item.get("source", ""))))
                cells[2].text = _xml_safe(parameter_label(str(item.get("name", ""))))
                cells[3].text = _xml_safe(str(item.get("raw", "")))
            _format_table(table, [2100, 1800, 2700, 2760])
        else:
            document.add_paragraph("Даты не извлечены.")

        document.add_heading("7. Выводы и заметки пользователя", level=1)
        document.add_heading("Выводы", level=2)
        document.add_paragraph(_xml_safe(evidence.user_conclusions or "Не заполнены."))
        document.add_heading("Заметки", level=2)
        document.add_paragraph(_xml_safe(evidence.user_notes or "Не заполнены."))

        document.add_heading("8. Ограничения", level=1)
        document.add_paragraph(
            "Программное обеспечение систематизирует технические факты и не устанавливает "
            "личность фактического автора, факт подделки или юридическую действительность подписи."
        )
        document.add_paragraph("Значения в таблицах длиннее 6000 символов сокращены с явной пометкой. Полные значения сохранены в базе и доступны в JSON-отчёте. Недопустимые в XML управляющие символы показаны как буквальные escape-последовательности.")
        if result.errors:
            for error in result.errors:
                document.add_paragraph(_xml_safe(error), style="List Bullet")

        document.add_page_break()
        document.add_heading("Приложение: необработанные результаты", level=1)
        for run in result.tool_runs:
            document.add_heading(_xml_safe(run.tool_name), level=2)
            self._key_value_table(
                document,
                [
                    ("Версия", run.version),
                    ("Команда", json.dumps(run.command, ensure_ascii=False)),
                    ("Статус", display_value(run.status.value)),
                    ("Exit code", str(run.exit_code)),
                ],
            )
            streams = [("stdout", run.stdout), ("stderr", run.stderr)]
            if run.parsed_result.get("stdin_utf8"):
                streams.insert(0, ("stdin (UTF-8)", str(run.parsed_result["stdin_utf8"])))
            for label, value in streams:
                document.add_heading(label, level=3)
                _add_preformatted(document, value or "(пусто)")

    def _key_value_table(self, document: DocumentType, rows: list[tuple[str, str]]) -> None:
        table = document.add_table(rows=0, cols=2)
        for label, value in rows:
            cells = table.add_row().cells
            cells[0].text = _xml_safe(str(label))
            cells[1].text = _xml_safe(str(value))
            for run in cells[0].paragraphs[0].runs:
                run.bold = True
        _format_table(table, [2700, 6660], header=False)

    def _tool_table(self, document: DocumentType, result: AnalysisResult) -> None:
        table = document.add_table(rows=1, cols=5)
        for cell, value in zip(
            table.rows[0].cells, ("Инструмент", "Версия", "Статус", "Код", "Длительность")
        ):
            cell.text = value
        for run in result.tool_runs:
            cells = table.add_row().cells
            for cell, value in zip(
                cells,
                (
                    run.tool_name,
                    run.version,
                    display_value(run.status.value),
                    str(run.exit_code),
                    f"{run.duration_ms} мс",
                ),
            ):
                cell.text = _xml_safe(value)
        _format_table(table, [1800, 2700, 1500, 900, 2460])

    def _provenance_table(
        self, document: DocumentType, result: AnalysisResult
    ) -> None:
        provenance = result.artifacts.get("pdf_provenance", {})
        if not isinstance(provenance, dict):
            document.add_paragraph("Для этого формата данные не формируются.")
            return
        caution = document.add_paragraph(_xml_safe(str(provenance.get("caution", ""))))
        if caution.runs:
            _font(caution.runs[0].font, "Calibri", 10, color=MUTED, bold=True)
        signals = provenance.get("signals", [])
        if not isinstance(signals, list) or not signals:
            document.add_paragraph("Технические признаки не извлечены.")
            return
        table = document.add_table(rows=1, cols=5)
        for cell, value in zip(
            table.rows[0].cells,
            ("Признак", "Значение", "Источник", "Сила признака", "Как трактовать"),
        ):
            cell.text = value
        for item in signals:
            if not isinstance(item, dict):
                continue
            cells = table.add_row().cells
            values = (
                item.get("role", ""),
                _table_value(_display(item.get("value", ""))),
                source_label(str(item.get("source", ""))),
                item.get("strength", ""),
                item.get("explanation", ""),
            )
            for cell, value in zip(cells, values):
                cell.text = _xml_safe(str(value))
        _format_table(table, [1700, 2500, 1500, 1600, 2060])

    def _finding_table(self, document: DocumentType, findings: Iterable[Finding]) -> None:
        table = document.add_table(rows=1, cols=5)
        for cell, value in zip(
            table.rows[0].cells,
            ("Параметр", "Значение", "Источник", "Оценка", "Что означает"),
        ):
            cell.text = value
        for item in findings:
            cells = table.add_row().cells
            values = (
                parameter_label(item.name),
                _table_value(display_value(item.value)),
                source_label(item.source) + ("\n" + item.location if item.location else ""),
                display_value(item.level.value) + "\nЭксперт: " + display_value(item.user_mark),
                parameter_description(item.name, item.details or item.description),
            )
            for cell, value in zip(cells, values):
                cell.text = _xml_safe(value)
            if item.conflict:
                for cell in cells:
                    _cell_fill(cell, "FFF1C9")
        _format_table(table, [1800, 2600, 1500, 1160, 2300])


def _table_value(value: str) -> str:
    if len(value) <= 6000:
        return value
    return value[:6000] + f"\n[Сокращено: всего {len(value)} символов. Полное значение — в JSON-отчёте.]"


def _format_table(table, widths: list[int], *, header: bool = True) -> None:
    table.autofit = False
    table.style = "Table Grid"
    table_element = table._tbl
    properties = table_element.tblPr
    width = properties.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(sum(widths)))
    indent = properties.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), "120")
    grid = table_element.tblGrid
    for child in list(grid):
        grid.remove(child)
    for item in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(item))
        grid.append(column)
    for row_index, row in enumerate(table.rows):
        row_properties = row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        if header and row_index == 0:
            row_properties.append(OxmlElement("w:tblHeader"))
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_width.type = "dxa"
            tc_width.w = widths[index]
            _cell_margins(cell, 80, 120, 80, 120)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    _font(run.font, "Calibri", 9.5, bold=((header and row_index == 0) or (not header and index == 0)))
            if header and row_index == 0:
                _cell_fill(cell, "F2F4F7")


def _cell_margins(cell, top: int, start: int, bottom: int, end: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _cell_fill(cell, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def _font(font, name: str, size: float, color: RGBColor | None = None, bold: bool = False) -> None:
    font.name = name
    font.size = Pt(size)
    font.bold = bold
    if color:
        font.color.rgb = color


def _add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def _add_preformatted(document: DocumentType, value: str) -> None:
    value = _xml_safe(value)
    for start in range(0, len(value), 20_000):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(value[start : start + 20_000])
        _font(run.font, "Consolas", 8)


def _xml_safe(value: str) -> str:
    """Preserve forbidden XML characters as readable escapes in DOCX output."""
    output: list[str] = []
    for character in value:
        codepoint = ord(character)
        allowed = (
            codepoint in (0x09, 0x0A, 0x0D)
            or 0x20 <= codepoint <= 0xD7FF
            or 0xE000 <= codepoint <= 0xFFFD
            or 0x10000 <= codepoint <= 0x10FFFF
        )
        if allowed:
            output.append(character)
        elif codepoint <= 0xFF:
            output.append(f"\\x{codepoint:02X}")
        else:
            output.append(f"\\u{codepoint:04X}")
    return "".join(output)


def _display(value: object) -> str:
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False, indent=2, default=str)
    return str(value)
