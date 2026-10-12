"""Capture actual Qt widgets using explicitly synthetic files, never user evidence."""
from __future__ import annotations

import json
import os
import sys
import zipfile
from datetime import datetime, UTC
from pathlib import Path

os.environ["QT_QPA_PLATFORM"] = "offscreen"
os.environ["QT_SCALE_FACTOR"] = "1"
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from PySide6.QtCore import QSettings, QPoint, QRect
from PySide6.QtGui import QFontDatabase
from PySide6.QtWidgets import QApplication

from app.core.config import AppConfig
from app.core.service import ForensicsService
from app.gui.dialogs import ProjectDialog
from app.gui.main_window import MainWindow
from app.models import AnalysisMode, ResearchProject
from app.resources.theme import apply_theme

GUIDE = ROOT / "docs/user-guide"
IMAGES = GUIDE / "screenshots"
EXAMPLES = GUIDE / "examples"
WORK = ROOT / "tmp/pdfs/guide-capture"


def create_examples():
    EXAMPLES.mkdir(parents=True, exist_ok=True)
    path = EXAMPLES / "Учебный договор.docx"
    doc = Document()
    doc.add_heading("Учебный договор", 0)
    doc.add_paragraph("Синтетический материал для знакомства с программой. Это не реальный договор.")
    paragraph = doc.add_paragraph("Условие оплаты: ")
    change = OxmlElement("w:ins")
    for name, value in {"id": "1", "author": "Учебный редактор", "date": "2026-02-02T09:00:00Z"}.items():
        change.set(qn(f"w:{name}"), value)
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "в течение 10 дней"
    run.append(text)
    change.append(run)
    paragraph._p.append(change)
    props = doc.core_properties
    props.title = "Учебный договор"
    props.author = "Учебный автор"
    props.last_modified_by = "Учебный редактор"
    props.created = datetime(2026, 2, 1, 7, tzinfo=UTC)
    props.modified = datetime(2026, 2, 1, 7, tzinfo=UTC)
    doc.save(path)
    with zipfile.ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for info, data in entries:
            if info.filename == "docProps/core.xml":
                data = data.replace(b'<dcterms:modified xsi:type="dcterms:W3CDTF">2026-02-01T07:00:00Z', b'<dcterms:modified xsi:type="dcterms:W3CDTF">2026-02-01T12:00:00+05:00')
            archive.writestr(info, data)
    exif = Image.Exif()
    exif[271], exif[272], exif[305] = "Training", "Synthetic image", "Training fixture"
    exif[306] = "2026:02:01 12:00:00"
    Image.new("RGB", (640, 400), "#7c83b8").save(EXAMPLES / "Учебное изображение.jpg", exif=exif)
    (EXAMPLES / "ПРОЧИТАЙТЕ.txt").write_text(
        "СИНТЕТИЧЕСКИЕ УЧЕБНЫЕ ФАЙЛЫ\n\n"
        "Учебный договор.docx: условные автор и редактор, одна сохранённая вставка.\n"
        "Даты created 2026-02-01T07:00:00Z и modified 2026-02-01T12:00:00+05:00\n"
        "обозначают один записанный момент. Дата вставки: 2026-02-02T09:00:00Z.\n"
        "Учебное изображение.jpg: однотонный прямоугольник; EXIF DateTime без пояса.\n"
        "Все сведения заданы генератором. Это не доказательства реальных событий.\n"
        "Размеры, хеши и даты файловой системы могут отличаться при повторном создании.\n",
        encoding="utf-8-sig",
    )
    return path


def main():
    IMAGES.mkdir(parents=True, exist_ok=True)
    WORK.mkdir(parents=True, exist_ok=True)
    QSettings.setDefaultFormat(QSettings.Format.IniFormat)
    QSettings.setPath(QSettings.Format.IniFormat, QSettings.Scope.UserScope, str(WORK / "qt-settings"))
    app = QApplication.instance() or QApplication([])
    for name in ("segoeui.ttf", "segoeuib.ttf", "seguisb.ttf", "consola.ttf"):
        QFontDatabase.addApplicationFont(str(Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / name))
    app.setStyle("Fusion")
    apply_theme(app)
    data = WORK / datetime.now().strftime("workspace-%Y%m%d-%H%M%S")
    config = AppConfig(data_dir=data, database_path=data / "db.sqlite3", cases_dir=data / "cases", reports_dir=data / "reports", validate_signatures=False)
    service = ForensicsService(config)
    window = MainWindow(service)
    window.resize(1366, 900)
    window.sidebar_action.setChecked(True)
    window.show()
    manifest = []

    def shot(name, widget=None, rect=None):
        app.processEvents()
        widget = widget or window
        pixmap = widget.grab(rect) if rect is not None else widget.grab()
        target = IMAGES / f"{name}.png"
        assert pixmap.save(str(target))
        manifest.append({"image": target.name, "width": pixmap.width(), "height": pixmap.height(), "source": "Actual Qt widget; synthetic training material"})

    shot("01-start")
    path = create_examples()
    project = ResearchProject("Учебное исследование", "УЧ-001", "Синтетический договор и изображение", source_description="Учебные файлы из комплекта программы")
    service.create_project(project)
    dialog = ProjectDialog(window, project)
    dialog.resize(780, 600)
    dialog.show()
    shot("02-project", dialog)
    dialog.close()
    result = service.analyze_file(path, project_id=project.id, mode=AnalysisMode.CUSTOM)
    service.analyze_file(EXAMPLES / "Учебное изображение.jpg", project_id=project.id, mode=AnalysisMode.CUSTOM)
    window._refresh_projects(project.id)
    window._refresh_files(result.evidence.id)
    window._display_result(result)
    window.mode_combo.setCurrentIndex(2)
    window.tabs.setCurrentIndex(0)
    shot("03-overview")
    assistant = window.research_assistant
    row = next(i for i, answer in enumerate(assistant.answers) if answer.rule_id == "authorship")
    assistant.table.setCurrentCell(row, 0)
    shot("04-assistant", assistant)
    assistant._open_ground()
    explorer = window.findings_explorer
    shot("05-finding", explorer)
    explorer.search.setText("нет_такой_записи")
    shot("06-empty-search", explorer, QRect(0, 0, explorer.width(), 215))
    explorer.reset_filters.click()
    creator = next(f for f in result.findings if f.name == "creator")
    explorer.reveal(creator.id)
    explorer.raw_toggle.setChecked(True)
    shot("07-raw", explorer)
    shot("07b-raw-value", explorer.raw_detail)
    explorer.raw_toggle.setChecked(False)
    explorer._add_observation()
    panel = window.chronology_panel
    window.tabs.setCurrentWidget(panel)
    for field, action in (("created", panel.pin_date), ("modified", panel.compare_dates)):
        finding = next(f for f in result.findings if f.name == field)
        row = next(i for i, event in enumerate(panel.events) if event.get("finding_id") == finding.id)
        panel.table.selectRow(row)
        action.click()
    assert "одному моменту" in panel.detail.toPlainText()
    shot("08-chronology", panel)
    panel._add_note()
    window.notes_edit.moveCursor(window.notes_edit.textCursor().MoveOperation.Start)
    shot("09-notes", window.conclusions_page)
    menu = window.save_report_button.menu()
    menu.popup(window.save_report_button.mapToGlobal(QPoint(0, window.save_report_button.height())))
    shot("10-export-menu", menu)
    menu.close()
    window._show_queue()
    window.queue_panel.queue.enqueue(path, project_id=project.id, mode=AnalysisMode.QUICK)
    window.queue_panel.refresh()
    shot("11-queue", window.queue_panel, QRect(0, 0, window.queue_panel.width(), 320))
    import tempfile
    with tempfile.TemporaryDirectory(prefix="DFI-training-") as temporary_watch:
        window.queue_panel.queue.save_watch(Path(temporary_watch), project.id, AnalysisMode.QUICK, False)
        window.queue_panel.refresh()
        window.queue_panel.pages.setCurrentIndex(1)
        shot("12-watch", window.queue_panel, QRect(0, 0, window.queue_panel.width(), 300))
    image_result = next(service.repository.load(item["id"]) for item in service.repository.list_evidence_summaries(project.id) if item["original_name"].endswith(".jpg"))
    window._display_result(image_result)
    window.tabs.setCurrentWidget(panel)
    row = next(i for i, event in enumerate(panel.events) if event["scope"] == "embedded_metadata" and not event.get("timezone_present"))
    panel.table.selectRow(row)
    shot("14-unknown-timezone", panel)
    window._display_result(result)
    window.tabs.setCurrentIndex(0)
    window.resize(900, 700)
    shot("13-small-window")
    window.close()
    (GUIDE / "screenshots.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"screenshots": len(manifest), "directory": str(IMAGES), "minimum_size": [window.minimumSizeHint().width(), window.minimumSizeHint().height()]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
